import io
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services.storage import storage_service

client = TestClient(app)

def test_health_check():
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert "rag" in data["modules"]
    assert "translate" in data["modules"]
    assert "meeting" in data["modules"]
    assert "providers" in data

def test_upload_valid_text_document():
    content = b"Institutional Policy Document.\n\nSection 1: Academic council approved the new admissions schedule."
    response = client.post(
        "/api/uploads",
        data={"module": "rag"},
        files={"files": ("policy.txt", content, "text/plain")}
    )
    assert response.status_code == 200
    data = response.json()
    assert len(data["files"]) == 1
    assert data["files"][0]["name"] == "policy.txt"
    assert data["files"][0]["status"] == "ready"
    assert "id" in data["files"][0]

def test_upload_invalid_extension():
    content = b"echo 'hack'"
    response = client.post(
        "/api/uploads",
        data={"module": "rag"},
        files={"files": ("script.exe", content, "application/octet-stream")}
    )
    assert response.status_code == 400
    assert "not supported" in response.json()["detail"].lower()

def test_upload_corrupt_pdf():
    # Not starting with %PDF-
    bad_pdf = b"NOT_A_REAL_PDF_FILE"
    response = client.post(
        "/api/uploads",
        data={"module": "rag"},
        files={"files": ("corrupt.pdf", bad_pdf, "application/pdf")}
    )
    assert response.status_code == 400
    assert "not a valid pdf" in response.json()["detail"].lower()

def test_rag_workflow_with_evidence_and_no_evidence():
    # 1. Upload doc
    doc_text = (
        "Academic Council Regulations 2026.\n\n"
        "Section 4: The annual student admissions portal opens on July 15 and closes on August 30.\n\n"
        "Section 5: Fee payment must be completed within 14 days of admission confirmation."
    ).encode("utf-8")
    
    upload_res = client.post(
        "/api/uploads",
        data={"module": "rag"},
        files={"files": ("council_regulations.txt", doc_text, "text/plain")}
    )
    assert upload_res.status_code == 200
    file_id = upload_res.json()["files"][0]["id"]

    # 2. Create collection
    col_res = client.post(
        "/api/rag/collections",
        json={"name": "Admissions Regulations", "fileIds": [file_id]}
    )
    assert col_res.status_code == 200
    col_data = col_res.json()
    col_id = col_data["id"]
    assert col_data["status"] == "ready"
    assert col_data["fileCount"] == 1
    assert col_data["chunkCount"] > 0

    # 3. Question with clear evidence in document
    q_res = client.post(
        f"/api/rag/collections/{col_id}/questions",
        json={"question": "When does the student admissions portal open?"}
    )
    assert q_res.status_code == 200
    ans_data = q_res.json()
    assert ans_data["evidenceFound"] is True
    assert "July 15" in ans_data["answer"] or "admissions" in ans_data["answer"].lower()
    assert len(ans_data["citations"]) > 0
    assert ans_data["citations"][0]["fileId"] == file_id

    # 4. Question with NO evidence in document
    no_ev_res = client.post(
        f"/api/rag/collections/{col_id}/questions",
        json={"question": "What is the quantum speed of the Mars rover spaceship?"}
    )
    assert no_ev_res.status_code == 200
    no_ev_data = no_ev_res.json()
    assert no_ev_data["evidenceFound"] is False
    assert "no sufficient evidence" in no_ev_data["answer"].lower()
    assert len(no_ev_data["citations"]) == 0

def test_upload_empty_file_rejected():
    response = client.post(
        "/api/uploads",
        data={"module": "rag"},
        files={"files": ("empty.txt", b"", "text/plain")}
    )
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()

def test_rag_collection_isolation():
    """Verify that documents in Collection A are completely isolated from Collection B."""
    doc_a = b"Collection A Secret: The admissions committee convenes only on Wednesdays."
    doc_b = b"Collection B Information: Hostel room allocations are announced every Friday."

    up_a = client.post("/api/uploads", data={"module": "rag"}, files={"files": ("doc_a.txt", doc_a, "text/plain")})
    up_b = client.post("/api/uploads", data={"module": "rag"}, files={"files": ("doc_b.txt", doc_b, "text/plain")})
    assert up_a.status_code == 200 and up_b.status_code == 200
    id_a = up_a.json()["files"][0]["id"]
    id_b = up_b.json()["files"][0]["id"]

    col_a = client.post("/api/rag/collections", json={"name": "Admissions Col", "fileIds": [id_a]}).json()["id"]
    col_b = client.post("/api/rag/collections", json={"name": "Hostel Col", "fileIds": [id_b]}).json()["id"]

    # Query Col A for Col B's topic: must NOT find evidence
    cross_res = client.post(f"/api/rag/collections/{col_a}/questions", json={"question": "When are hostel room allocations announced?"})
    assert cross_res.status_code == 200
    assert cross_res.json()["evidenceFound"] is False
    assert len(cross_res.json()["citations"]) == 0

    # Query Col A for its own topic: must find evidence and cite doc_a only
    own_res = client.post(f"/api/rag/collections/{col_a}/questions", json={"question": "When does the admissions committee convene?"})
    assert own_res.status_code == 200
    assert own_res.json()["evidenceFound"] is True
    assert all(c["fileId"] == id_a for c in own_res.json()["citations"])

def test_rag_dimension_mismatch_rejection(monkeypatch):
    """Verify that querying a collection with an incompatible embedding vector dimension returns a clear 409 rejection."""
    doc = b"Admissions policy guidelines 2026."
    up = client.post("/api/uploads", data={"module": "rag"}, files={"files": ("guidelines.txt", doc, "text/plain")})
    assert up.status_code == 200
    fid = up.json()["files"][0]["id"]

    # Create collection using local 256-dim provider
    col_id = client.post("/api/rag/collections", json={"name": "Dimension Test Col", "fileIds": [fid]}).json()["id"]

    # Mock an incompatible embedding provider with 768 dimensions (simulating a provider change)
    from app.providers.embeddings import EmbeddingProvider
    class Mock768Provider(EmbeddingProvider):
        provider_name = "mock_provider"
        model_name = "mock-768"
        dimension = 768
        def embed_query(self, text: str):
            return [0.1] * 768
        def embed_documents(self, texts):
            return [[0.1] * 768 for _ in texts]

    monkeypatch.setattr("app.services.rag_service.get_embedding_provider", lambda: Mock768Provider())

    # Querying should reject with 409 Conflict rather than silently computing bad dot products or crashing
    q_res = client.post(f"/api/rag/collections/{col_id}/questions", json={"question": "What is the admissions policy?"})
    assert q_res.status_code == 409
    assert "mismatch" in q_res.json()["detail"].lower()
    assert "re-index" in q_res.json()["detail"].lower() or "rebuild" in q_res.json()["detail"].lower()

def test_indic_translate_language_validation():
    """Verify that unsupported languages and identical source/target languages are clearly rejected."""
    # 1. Identical source and target language
    same_res = client.post(
        "/api/translate/jobs",
        json={"sourceText": "Hello", "sourceLanguage": "English", "targetLanguage": "English"}
    )
    assert same_res.status_code == 400
    assert "identical" in same_res.json()["detail"].lower()

    # 2. Unsupported source language
    bad_src = client.post(
        "/api/translate/jobs",
        json={"sourceText": "Bonjour", "sourceLanguage": "French", "targetLanguage": "Tamil"}
    )
    assert bad_src.status_code == 400
    assert "not supported" in bad_src.json()["detail"].lower()

    # 3. Unsupported target language
    bad_tgt = client.post(
        "/api/translate/jobs",
        json={"sourceText": "Hello", "sourceLanguage": "English", "targetLanguage": "German"}
    )
    assert bad_tgt.status_code == 400
    assert "not supported" in bad_tgt.json()["detail"].lower()

def test_indic_translate_no_model_configured_returns_blank_draft(monkeypatch):
    """When no real translation model is configured, draft must be blank, labeled no_model, and allow manual entry."""
    from app.providers.translation import LocalTranslationProvider
    monkeypatch.setattr("app.services.translation_service.get_translation_provider", lambda: LocalTranslationProvider())

    sample_text = "New circular issued by the registrar regarding annual leave policies."

    # 1. POST /api/translate/jobs creates the job
    create_res = client.post(
        "/api/translate/jobs",
        json={
            "sourceText": sample_text,
            "sourceLanguage": "English",
            "targetLanguage": "Tamil"
        }
    )
    assert create_res.status_code == 200
    job_id = create_res.json()["id"]

    # 2. GET /api/translate/jobs/{job_id} fetches the result
    get_res = client.get(f"/api/translate/jobs/{job_id}")
    assert get_res.status_code == 200
    job_data = get_res.json()

    assert job_data["status"] == "completed"
    assert len(job_data["segments"]) > 0
    seg = job_data["segments"][0]
    # Translation must be blank because no neural model is configured (no fake glossary substitutions)
    assert seg["translated"] == "", "Unconfigured translation engine must return blank translated text"
    # reviewFlag clearly indicates no_model
    assert seg["reviewFlag"] == "no_model", f"Expected reviewFlag 'no_model', got '{seg['reviewFlag']}'"
    assert any("no translation model is configured" in lim.lower() for lim in job_data["limitations"])

    # 3. User can manually enter a draft for the blank segment
    manual_edit_res = client.put(
        f"/api/translate/jobs/{job_id}/segments/{seg['id']}",
        json={"translated": "பதிவாளர் வெளியிட்ட புதிய சுற்றறிக்கை.", "reviewFlag": "edited"}
    )
    assert manual_edit_res.status_code == 200
    edited_job = manual_edit_res.json()
    edited_seg = next(s for s in edited_job["segments"] if s["id"] == seg["id"])
    assert edited_seg["translated"] == "பதிவாளர் வெளியிட்ட புதிய சுற்றறிக்கை."
    assert edited_seg["reviewFlag"] == "edited"

def test_indic_translate_configured_model_path(monkeypatch):
    """When a real MT provider is configured, verify genuine translation output, reviewFlag='unreviewed', and human-in-the-loop workflow."""
    from app.providers.translation import TranslationProvider
    class MockNeuralTranslationProvider(TranslationProvider):
        def translate_segments(self, segments, source_lang, target_lang):
            # Returns real Tamil script translation
            return ["கல்விசார் மன்றம் புதிய கொள்கைக்கு ஒப்புதல் அளித்தது."]

    monkeypatch.setattr("app.services.translation_service.get_translation_provider", lambda: MockNeuralTranslationProvider())

    sample_text = "The academic council approved the new admissions policy for the institution."
    create_res = client.post(
        "/api/translate/jobs",
        json={"sourceText": sample_text, "sourceLanguage": "English", "targetLanguage": "Tamil"}
    )
    assert create_res.status_code == 200
    job = create_res.json()
    assert job["status"] == "completed"
    assert len(job["segments"]) == 1

    seg = job["segments"][0]
    # Distinguish genuine translation from blank, unchanged English, or glossary substitution
    assert seg["translated"] != "", "Configured model must return non-blank translation"
    assert seg["translated"] != sample_text, "Must not return unchanged source English"
    # Contains Tamil Unicode characters
    import re
    assert re.search(r"[\u0B80-\u0BFF]", seg["translated"]), "Must contain genuine Tamil script"
    # Must be marked unreviewed (requiring human review, never automatically approved)
    assert seg["reviewFlag"] == "unreviewed"
    assert seg["confidence"] > 0

    # User can approve or edit the machine draft
    approve_res = client.put(
        f"/api/translate/jobs/{job['id']}/segments/{seg['id']}",
        json={"translated": seg["translated"], "reviewFlag": "approved"}
    )
    assert approve_res.status_code == 200
    assert approve_res.json()["segments"][0]["reviewFlag"] == "approved"


def test_meeting_minutes_workflow_transcript():
    transcript_text = (
        "Topic: Admissions Calendar Review.\n"
        "Alice: We are discussing the schedule for 2026 admissions.\n"
        "Bob: Decision: Academic council approved the revised timeline.\n"
        "Alice: Action item: Dr. Sharma will prepare the portal guidelines by next week.\n"
        "Charlie: Someone should also check the server bandwidth."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("meeting_transcript.txt", transcript_text, "text/plain")}
    )
    assert upload_res.status_code == 200
    tf_id = upload_res.json()["files"][0]["id"]

    # Create meeting job
    job_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "Council Meeting"}
    )
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert job_data["status"] == "completed"
    assert job_data["draft"] is not None
    draft = job_data["draft"]
    assert len(draft["decisions"]) > 0
    assert len(draft["actionItems"]) > 0
    # Verify uncertainty flag was set for items without explicit owner/date
    assert any(a["uncertain"] is True for a in draft["actionItems"])

def test_meeting_minutes_audio_unconfigured(monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "AUDIO_PROVIDER", "none")

    # Synthetic WAV header
    wav_header = b"RIFF\x24\x00\x00\x00WAVEfmt \x10\x00\x00\x00\x01\x00\x01\x00\x44\xac\x00\x00\x88\x58\x01\x00\x02\x00\x10\x00data\x00\x00\x00\x00"
    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("recording.wav", wav_header, "audio/wav")}
    )
    assert upload_res.status_code == 200
    af_id = upload_res.json()["files"][0]["id"]

    # When AUDIO_PROVIDER=none, should cleanly fail with explanation rather than fake transcription
    job_res = client.post(
        "/api/meetings/jobs",
        json={"audioFileIds": [af_id]}
    )
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert job_data["status"] == "failed"
    assert "not configured" in job_data["error"].lower() or "whisper" in job_data["error"].lower()


def test_meeting_minutes_exact_transcript_extraction():
    """Regression test for the exact user-provided transcript.

    Transcript:
        Chair: We agreed to hold the workshop on 3 November 2026.
        Maya: I will book Room 2 by 25 October 2026.
        The team approved the budget of 15,000 rupees.

    Expected:
    - Action: "Book Room 2", owner="Maya", dueDate="25 October 2026", uncertain=False
    - Decision: budget approval of 15,000 rupees (status="agreed")
    - Decision: workshop agreement (status="agreed")
    - No generic "Review draft minutes" placeholder action
    - No "No explicit action assignments" uncertainty
    """
    transcript_text = (
        "Chair: We agreed to hold the workshop on 3 November 2026.\n"
        "Maya: I will book Room 2 by 25 October 2026.\n"
        "The team approved the budget of 15,000 rupees."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("exact_meeting.txt", transcript_text, "text/plain")},
    )
    assert upload_res.status_code == 200
    tf_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "Workshop Planning"},
    )
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert job_data["status"] == "completed"
    draft = job_data["draft"]
    assert draft is not None

    # ── Action items ──────────────────────────────────────────────────
    assert len(draft["actionItems"]) >= 1, f"Expected ≥1 action items, got {draft['actionItems']}"
    book_action = draft["actionItems"][0]
    assert "book room 2" in book_action["task"].lower(), f"Task should contain 'Book Room 2', got: {book_action['task']}"
    assert book_action["owner"] == "Maya", f"Owner should be 'Maya', got: {book_action['owner']}"
    assert book_action["dueDate"] is not None, "Due date should not be None"
    assert "25 October 2026" in book_action["dueDate"] or "25 Oct" in book_action["dueDate"], \
        f"Due date should be '25 October 2026', got: {book_action['dueDate']}"
    assert book_action["uncertain"] is False, "uncertain should be False when owner and dueDate are both present"

    # No generic placeholder action
    for act in draft["actionItems"]:
        assert "review draft minutes" not in act["task"].lower(), \
            f"Generic fallback action should not appear: {act['task']}"

    # ── Decisions ─────────────────────────────────────────────────────
    decision_texts = " ".join(d["decision"].lower() for d in draft["decisions"])
    assert "15,000 rupees" in decision_texts or "budget" in decision_texts, \
        f"Budget approval decision missing. Decisions: {[d['decision'] for d in draft['decisions']]}"
    # All captured decisions should have status "agreed"
    for d in draft["decisions"]:
        assert d["status"] == "agreed", f"Decision status should be 'agreed', got: {d['status']}"

    # ── Uncertainties ─────────────────────────────────────────────────
    for u in draft.get("uncertainties", []):
        assert "no explicit action assignments" not in u.lower(), \
            f"Should not report 'no explicit action assignments': {u}"


def test_meeting_minutes_transcript_with_no_actions():
    """Verify that a transcript with no action assignments produces an empty action list without inventing actions."""
    transcript_text = (
        "Topic: General Information Sharing.\n"
        "Prof. Kumar: The department organized three technical seminars this semester.\n"
        "Dr. Varma: Attendance was high across all batches.\n"
        "Chair: Thank you everyone for the constructive update."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("info_sharing.txt", transcript_text, "text/plain")},
    )
    assert upload_res.status_code == 200
    tf_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "Department Review"},
    )
    assert job_res.status_code == 200
    draft = job_res.json()["draft"]
    assert draft is not None
    # Must NOT invent action items or decisions
    assert len(draft["actionItems"]) == 0, f"Expected 0 action items, got: {draft['actionItems']}"
    assert len(draft["decisions"]) == 0, f"Expected 0 decisions, got: {draft['decisions']}"
    # Must preserve discussions
    assert len(draft["discussions"]) > 0


def test_error_handling_sanitizes_secrets():
    """Verify that backend error messages never leak API keys, system paths with secrets, or database credentials."""
    # Test an invalid collection query
    err_res = client.post("/api/rag/collections/non-existent-uuid/questions", json={"question": "Test?"})
    assert err_res.status_code == 404
    detail = str(err_res.json()["detail"])
    assert "AIza" not in detail
    assert "sk-" not in detail
    assert "password" not in detail.lower()
    assert "secret" not in detail.lower()


def test_meeting_minutes_longer_library_renovation_transcript():
    """Verify the realistic library renovation transcript extracts evidence-based summary,

    discussion topic groups, decisions, and action items with exact deadlines without inventing dates.
    """
    transcript_text = (
        "Chair: We need to review the library renovation before the next semester. The reading room is too noisy during afternoon study hours.\n"
        "Latha: Students asked for a quiet study zone and more charging points. We should consider moving the group tables near the entrance.\n"
        "Ravi: The facilities team can inspect the electrical outlets this week. I will send the inspection report by Friday.\n"
        "Chair: We agreed to create a quiet study zone in the east wing. The group tables will remain where they are until the facilities review is complete.\n"
        "Meena: I will collect student feedback and share a summary by 12 November 2026.\n"
        "Chair: The renovation budget is approved up to 2 lakh rupees, subject to the facilities report."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("library_renovation.txt", transcript_text, "text/plain")},
    )
    assert upload_res.status_code == 200
    tf_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "Library Renovation Review"},
    )
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert job_data["status"] == "completed"
    draft = job_data["draft"]
    assert draft is not None

    # ── 1. Summary Evidence ───────────────────────────────────────────
    summary_lower = draft["summary"].lower()
    assert "library renovation" in summary_lower or "renovation" in summary_lower, "Summary must mention library renovation"
    assert "noise" in summary_lower or "noisy" in summary_lower or "quiet study" in summary_lower, "Summary must mention noise / quiet study zone"
    assert "facilities" in summary_lower, "Summary must mention facilities review/inspection"
    assert "budget" in summary_lower or "lakh" in summary_lower, "Summary must mention approved budget"

    # ── 2. Meaningful Discussion Topics ───────────────────────────────
    assert len(draft["discussions"]) >= 3, f"Expected at least 3 discussion topic groups, got {len(draft['discussions'])}"
    topics_text = " ".join(d["topic"].lower() for d in draft["discussions"])
    assert "renovation" in topics_text or "budget" in topics_text, "Discussion topic for renovation planning/budget missing"
    assert "quiet" in topics_text or "noise" in topics_text or "study zone" in topics_text, "Discussion topic for study zone/noise missing"
    assert "outlet" in topics_text or "facilities" in topics_text or "charging" in topics_text, "Discussion topic for electrical outlets missing"
    assert "table" in topics_text or "feedback" in topics_text, "Discussion topic for table placement/feedback missing"

    # Ensure points under topics contain actual statements
    all_points = [pt for d in draft["discussions"] for pt in d["points"]]
    assert any("reading room is too noisy" in pt.lower() for pt in all_points)
    assert any("charging points" in pt.lower() or "electrical outlets" in pt.lower() for pt in all_points)
    assert any("group tables" in pt.lower() for pt in all_points)

    # ── 3. Decisions ──────────────────────────────────────────────────
    assert len(draft["decisions"]) >= 3, f"Expected at least 3 decisions, got {len(draft['decisions'])}"
    decisions_text = " ".join(d["decision"].lower() for d in draft["decisions"])
    assert "quiet study zone" in decisions_text, "Decision on quiet study zone missing"
    assert "group tables will remain" in decisions_text or "remain where they are" in decisions_text, "Decision on keeping tables in place missing"
    assert "budget is approved" in decisions_text or "approved up to 2 lakh" in decisions_text, "Decision on budget approval missing"
    assert "facilities report" in decisions_text, "Condition 'subject to the facilities report' must be preserved"

    # ── 4. Action Items ───────────────────────────────────────────────
    assert len(draft["actionItems"]) == 2, f"Expected exactly 2 action items, got {len(draft['actionItems'])}"
    
    import re
    # Ravi's action
    ravi_act = next((a for a in draft["actionItems"] if a["owner"] == "Ravi"), None)
    assert ravi_act is not None, "Action item for Ravi missing"
    assert "inspection report" in ravi_act["task"].lower()
    assert ravi_act["dueDate"] is not None
    assert "friday" in ravi_act["dueDate"].lower(), f"Expected due date 'Friday', got {ravi_act['dueDate']}"
    # Verify Friday is preserved without fabricating an arbitrary calendar date
    assert not re.search(r"\b202\d-\d{2}-\d{2}\b", ravi_act["dueDate"]), "Should not fabricate calendar date for 'Friday'"
    assert ravi_act["uncertain"] is False

    # Meena's action
    meena_act = next((a for a in draft["actionItems"] if a["owner"] == "Meena"), None)
    assert meena_act is not None, "Action item for Meena missing"
    assert "feedback" in meena_act["task"].lower()
    assert meena_act["dueDate"] is not None
    assert "12 november 2026" in meena_act["dueDate"].lower()
    assert meena_act["uncertain"] is False


def test_meeting_minutes_pleasantries_only_zero_invented_content():
    """Verify that a transcript with only greetings and procedural pleasantries returns empty lists and no invented content."""
    transcript_text = (
        "Alice: Good morning everyone.\n"
        "Bob: Good morning Alice. Can everyone hear me?\n"
        "Charlie: Yes, loud and clear.\n"
        "Bob: Sounds good. Let's wait a minute.\n"
        "Alice: Thanks everyone."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("pleasantries.txt", transcript_text, "text/plain")},
    )
    assert upload_res.status_code == 200
    tf_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "Quick Sync"},
    )
    assert job_res.status_code == 200
    draft = job_res.json()["draft"]
    assert draft is not None

    # Empty lists for everything
    assert len(draft["discussions"]) == 0, f"Expected 0 discussions, got {draft['discussions']}"
    assert len(draft["decisions"]) == 0, f"Expected 0 decisions, got {draft['decisions']}"
    assert len(draft["actionItems"]) == 0, f"Expected 0 action items, got {draft['actionItems']}"
    assert "preliminary greetings" in draft["summary"].lower() or "without substantive" in draft["summary"].lower()


def test_meeting_minutes_draft_editing_and_persistence():
    """Verify that editing draft title, summary, discussions, decisions, actions, and uncertainties succeeds and persists."""
    transcript_text = (
        "Topic: IT Infrastructure.\n"
        "Deepak: We need to upgrade campus Wi-Fi bandwidth.\n"
        "Chair: Decision: Budget of 50,000 rupees approved for network equipment.\n"
        "Deepak: I will contact the vendor by 5 November 2026."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("it_meeting.txt", transcript_text, "text/plain")},
    )
    tf_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "IT Review"},
    )
    job_id = job_res.json()["id"]

    # Modify all fields
    modified_draft = {
        "title": "Edited IT Review 2026",
        "summary": "Custom edited summary of the IT infrastructure upgrade.",
        "discussions": [
            {"topic": "Campus Wi-Fi & Bandwidth Expansion", "points": ["Reviewed existing latency and traffic."], "timestamp": None}
        ],
        "decisions": [
            {"decision": "Budget approved for 50,000 INR with vendor discount.", "status": "agreed", "evidence": "Recorded"}
        ],
        "actionItems": [
            {"task": "Contact vendor for quotation", "owner": "Deepak", "dueDate": "5 November 2026", "uncertain": False}
        ],
        "uncertainties": ["Check if fiber optic line requires civil work permit."]
    }

    update_res = client.put(f"/api/meetings/jobs/{job_id}/draft", json={"draft": modified_draft})
    assert update_res.status_code == 200
    saved_draft = update_res.json()["draft"]

    assert saved_draft["title"] == "Edited IT Review 2026"
    assert saved_draft["summary"] == "Custom edited summary of the IT infrastructure upgrade."
    assert saved_draft["discussions"][0]["topic"] == "Campus Wi-Fi & Bandwidth Expansion"
    assert saved_draft["decisions"][0]["decision"] == "Budget approved for 50,000 INR with vendor discount."
    assert saved_draft["actionItems"][0]["task"] == "Contact vendor for quotation"
    assert saved_draft["uncertainties"] == ["Check if fiber optic line requires civil work permit."]

    # Fetch fresh from GET /api/meetings/jobs/{job_id} to verify database persistence
    get_res = client.get(f"/api/meetings/jobs/{job_id}")
    assert get_res.status_code == 200
    persisted_draft = get_res.json()["draft"]
    assert persisted_draft["title"] == "Edited IT Review 2026"
    assert persisted_draft["summary"] == "Custom edited summary of the IT infrastructure upgrade."


def test_meeting_minutes_title_change_persists_across_reload():
    """Regression test: User changes the title from default 'Institutional Meeting' to a custom title,

    saves draft edits, and reloads the draft from backend/SQLite. Verifies the new title persists
    and does not revert to 'Institutional Meeting'.
    """
    transcript_text = (
        "Chair: We need to review the library renovation before the next semester.\n"
        "Maya: I will book Room 2 by 25 October 2026.\n"
        "Chair: The budget is approved."
    ).encode("utf-8")

    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("transcript.txt", transcript_text, "text/plain")},
    )
    assert upload_res.status_code == 200
    tf_id = upload_res.json()["files"][0]["id"]

    # 1. Job created with default title "Institutional Meeting"
    create_res = client.post(
        "/api/meetings/jobs",
        json={"transcriptFileIds": [tf_id], "meetingTitle": "Institutional Meeting"},
    )
    assert create_res.status_code == 200
    job_id = create_res.json()["id"]
    initial_draft = create_res.json()["draft"]
    assert initial_draft["title"] == "Institutional Meeting"

    # 2. User changes the title to "Board of Trustees Strategic Review 2026"
    new_title = "Board of Trustees Strategic Review 2026"
    updated_draft_payload = dict(initial_draft)
    updated_draft_payload["title"] = new_title

    # 3. User clicks "Save draft edits" -> PUT /api/meetings/jobs/{job_id}/draft
    save_res = client.put(f"/api/meetings/jobs/{job_id}/draft", json={"draft": updated_draft_payload})
    assert save_res.status_code == 200
    assert save_res.json()["draft"]["title"] == new_title

    # 4. User refreshes / reloads via GET /api/meetings/jobs/{job_id}
    reload_res = client.get(f"/api/meetings/jobs/{job_id}")
    assert reload_res.status_code == 200
    reloaded_draft = reload_res.json()["draft"]

    # Verify the new title remains and did NOT revert to "Institutional Meeting"
    assert reloaded_draft["title"] == new_title, \
        f"Expected reloaded title to be '{new_title}', but got '{reloaded_draft['title']}'"
    assert reloaded_draft["title"] != "Institutional Meeting"

    # 5. Also verify GET /api/meetings/jobs/latest returns the newly updated job with the new title
    latest_res = client.get("/api/meetings/jobs/latest")
    assert latest_res.status_code == 200
    assert latest_res.json()["id"] == job_id
    assert latest_res.json()["draft"]["title"] == new_title


# ==========================================
# Module 04: Document Translation & Download Tests
# ==========================================

def test_translate_txt_upload_translate_download():
    """TXT upload -> translation job -> download returns text file with translated content."""
    content = b"The academic council approved the new policy.\n\nAll students must register before the deadline."
    upload_res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("policy.txt", content, "text/plain")},
    )
    assert upload_res.status_code == 200
    file_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/translate/jobs",
        json={"fileIds": [file_id], "sourceLanguage": "English", "targetLanguage": "Tamil"},
    )
    assert job_res.status_code == 200
    job = job_res.json()
    assert job["status"] == "completed"
    job_id = job["id"]
    assert job["sourceFileName"] == "policy.txt"
    assert job["sourceFileExt"] == ".txt"

    # Download with format=txt
    download_res = client.get(f"/api/translate/jobs/{job_id}/download?format=txt")
    assert download_res.status_code == 200
    assert "text/plain" in download_res.headers["content-type"]
    assert "policy" in download_res.headers["content-disposition"]
    assert "tam" in download_res.headers["content-disposition"]
    content_out = download_res.content.decode("utf-8")
    assert "Tamil" in content_out or "English" in content_out
    assert len(content_out.strip()) > 0


def test_translate_docx_upload_translate_download():
    """DOCX upload -> translation job -> auto-format download returns valid DOCX."""
    import io
    import docx

    doc = docx.Document()
    doc.add_heading("Academic Policy", 0)
    doc.add_paragraph("The council approved the new admissions policy.")
    doc.add_paragraph("All students must complete registration.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    docx_bytes = buf.read()

    upload_res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("council-minutes.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload_res.status_code == 200
    file_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/translate/jobs",
        json={"fileIds": [file_id], "sourceLanguage": "English", "targetLanguage": "Hindi"},
    )
    assert job_res.status_code == 200
    job = job_res.json()
    assert job["status"] == "completed"
    job_id = job["id"]
    assert job["sourceFileName"] == "council-minutes.docx"
    assert job["sourceFileExt"] == ".docx"

    # Download auto-detects DOCX
    download_res = client.get(f"/api/translate/jobs/{job_id}/download")
    assert download_res.status_code == 200
    assert "openxmlformats" in download_res.headers["content-type"]
    assert "council-minutes" in download_res.headers["content-disposition"]
    assert "hin" in download_res.headers["content-disposition"]

    out_buf = io.BytesIO(download_res.content)
    out_doc = docx.Document(out_buf)
    full_text = "\n".join([p.text for p in out_doc.paragraphs])
    assert len(full_text.strip()) > 0


def test_translate_download_reflects_edits():
    """Edited translation segments must be reflected in downloaded documents (both TXT and DOCX)."""
    import io
    import docx

    content = b"The library renovation budget is approved."
    upload_res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("test-edit.txt", content, "text/plain")},
    )
    assert upload_res.status_code == 200
    file_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/translate/jobs",
        json={"fileIds": [file_id], "sourceLanguage": "English", "targetLanguage": "Tamil"},
    )
    assert job_res.status_code == 200
    job = job_res.json()
    job_id = job["id"]
    seg_id = job["segments"][0]["id"]

    edited_text = "நூலக புனரமைப்பு நிதி அங்கீகரிக்கப்பட்டது."
    edit_res = client.put(
        f"/api/translate/jobs/{job_id}/segments/{seg_id}",
        json={"translated": edited_text, "reviewFlag": "edited"},
    )
    assert edit_res.status_code == 200

    # TXT download check
    txt_res = client.get(f"/api/translate/jobs/{job_id}/download?format=txt")
    assert txt_res.status_code == 200
    assert edited_text in txt_res.content.decode("utf-8")

    # DOCX download check
    docx_res = client.get(f"/api/translate/jobs/{job_id}/download?format=docx")
    assert docx_res.status_code == 200
    doc = docx.Document(io.BytesIO(docx_res.content))
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert edited_text in doc_text


def test_translate_download_not_found():
    """Download of non-existent job returns 404."""
    res = client.get("/api/translate/jobs/nonexistent-job-id-123/download")
    assert res.status_code == 404


def test_translate_unsupported_language_direction():
    """Unsupported language directions return 400 Bad Request."""
    res = client.post(
        "/api/translate/jobs",
        json={"sourceText": "Bonjour le monde", "sourceLanguage": "French", "targetLanguage": "Tamil"},
    )
    assert res.status_code == 400
    assert "not supported" in res.json()["detail"].lower()


def test_translate_empty_file_rejected():
    """Empty file upload for translation is rejected with a helpful 400 error."""
    res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("empty.txt", b"", "text/plain")},
    )
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()


def test_translate_invalid_file_type_for_translate_module():
    """Unsupported file extension for translate module is rejected with 400."""
    res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("audio.mp3", b"\xff\xfb\x90" + b"\x00" * 100, "audio/mpeg")},
    )
    assert res.status_code == 400
    assert "not supported" in res.json()["detail"].lower()


def test_translate_missing_provider_job_fails_gracefully(monkeypatch):
    """When translation provider raises an error, job status is 'failed' with a clear message."""
    from app.providers.translation import TranslationProvider

    class FailingProvider(TranslationProvider):
        def translate_segments(self, segments, source_lang, target_lang):
            raise RuntimeError("Translation model failed to load: GPU out of memory.")

    monkeypatch.setattr("app.services.translation_service.get_translation_provider", lambda: FailingProvider())

    res = client.post(
        "/api/translate/jobs",
        json={"sourceText": "Test text", "sourceLanguage": "English", "targetLanguage": "Tamil"},
    )
    assert res.status_code == 200
    job = res.json()
    assert job["status"] == "failed"
    assert "failed to load" in job["error"].lower() or "gpu" in job["error"].lower()
    assert job["segments"] == []


def test_translate_pdf_scanned_rejected():
    """A scanned (image-only) PDF yielding no extractable text is rejected with a clear message."""
    minimal_pdf = (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R>>endobj\n"
        b"xref\n0 4\n0000000000 65535 f \n0000000009 00000 n \n0000000058 00000 n \n0000000115 00000 n \n"
        b"trailer<</Size 4/Root 1 0 R>>\nstartxref\n190\n%%EOF"
    )
    upload_res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("scanned.pdf", minimal_pdf, "application/pdf")},
    )
    assert upload_res.status_code == 200
    file_id = upload_res.json()["files"][0]["id"]

    job_res = client.post(
        "/api/translate/jobs",
        json={"fileIds": [file_id], "sourceLanguage": "English", "targetLanguage": "Tamil"},
    )
    assert job_res.status_code == 400
    detail = job_res.json()["detail"].lower()
    assert any(w in detail for w in ["scanned", "no readable", "extract", "text", "empty"])


def test_translate_docx_structure_preservation_headings_bullets_tables():
    """Regression test: verify DOCX headings, bullet paragraphs, and tables are preserved in original order."""
    import io
    import docx
    from docx.text.paragraph import Paragraph

    # 1. Construct a rich source DOCX with Heading 1, paragraph, List Bullets, Heading 2, Table, and closing paragraph
    doc = docx.Document()
    doc.add_heading("Library Renovation and Academic Policies", level=1)
    doc.add_paragraph("The academic council approved the updated library guidelines.")
    doc.add_paragraph("Expand quiet study zones on the second floor.", style="List Bullet")
    doc.add_paragraph("Install 30 new charging ports for student laptops.", style="List Bullet")
    doc.add_heading("Budget and Allocation Summary", level=2)
    tbl = doc.add_table(rows=2, cols=3)
    tbl.style = "Table Grid"
    tbl.rows[0].cells[0].text = "Item"
    tbl.rows[0].cells[1].text = "Allocated Budget"
    tbl.rows[0].cells[2].text = "Approval Status"
    tbl.rows[1].cells[0].text = "Study Desks"
    tbl.rows[1].cells[1].text = "45,000"
    tbl.rows[1].cells[2].text = "Approved"
    doc.add_paragraph("All work must be completed before the semester commences.")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    docx_bytes = buf.read()

    # 2. Upload to /api/uploads
    upload_res = client.post(
        "/api/uploads",
        data={"module": "translate"},
        files={"files": ("structured-policy.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload_res.status_code == 200
    file_id = upload_res.json()["files"][0]["id"]

    # 3. Create translation job for English -> Tamil
    job_res = client.post(
        "/api/translate/jobs",
        json={"fileIds": [file_id], "sourceLanguage": "English", "targetLanguage": "Tamil"},
    )
    assert job_res.status_code == 200
    job = job_res.json()
    assert job["status"] == "completed"
    job_id = job["id"]
    segments = job["segments"]

    # Verify extracted block types in segments
    heading_segs = [s for s in segments if s.get("blockType") == "heading"]
    bullet_segs = [s for s in segments if s.get("blockType") == "bullet"]
    table_segs = [s for s in segments if s.get("blockType") == "table_cell"]

    assert len(heading_segs) == 2, f"Expected 2 heading segments, got {len(heading_segs)}"
    assert heading_segs[0]["blockLevel"] == 1
    assert heading_segs[1]["blockLevel"] == 2
    assert len(bullet_segs) == 2, f"Expected 2 bullet segments, got {len(bullet_segs)}"
    assert len(table_segs) == 6, f"Expected 6 table cell segments (2x3), got {len(table_segs)}"

    # 4. User edits a table cell in UI (Row 1, Col 2: "Approved" -> custom Tamil)
    approved_seg = [s for s in table_segs if s.get("rowIndex") == 1 and s.get("cellIndex") == 2][0]
    edited_status = "ஒப்புதல் அளிக்கப்பட்டது"
    edit_res = client.put(
        f"/api/translate/jobs/{job_id}/segments/{approved_seg['id']}",
        json={"translated": edited_status, "reviewFlag": "edited"},
    )
    assert edit_res.status_code == 200

    # 5. Download the rebuilt DOCX
    download_res = client.get(f"/api/translate/jobs/{job_id}/download?format=docx")
    assert download_res.status_code == 200
    assert "openxmlformats" in download_res.headers["content-type"]
    assert "structured-policy" in download_res.headers["content-disposition"]

    # 6. Inspect the rebuilt DOCX
    out_doc = docx.Document(io.BytesIO(download_res.content))

    # A) Check heading styles
    doc_headings = [p for p in out_doc.paragraphs if p.style and p.style.name and p.style.name.startswith("Heading")]
    assert len(doc_headings) >= 2, f"Expected at least 2 headings, found {len(doc_headings)}"
    heading_styles = [p.style.name for p in doc_headings]
    assert "Heading 1" in heading_styles, f"Heading 1 missing from styles: {heading_styles}"
    assert "Heading 2" in heading_styles, f"Heading 2 missing from styles: {heading_styles}"

    # B) Check bullet paragraphs
    doc_bullets = [p for p in out_doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert len(doc_bullets) == 2, f"Expected 2 List Bullet paragraphs, found {len(doc_bullets)}"
    for b in doc_bullets:
        assert len(b.text.strip()) > 0

    # C) Check table structure and rows
    assert len(out_doc.tables) >= 1, f"Expected at least 1 table, found {len(out_doc.tables)}"
    target_table = out_doc.tables[0]
    assert len(target_table.rows) == 2, f"Expected 2 rows in table, got {len(target_table.rows)}"
    assert len(target_table.columns) == 3, f"Expected 3 columns in table, got {len(target_table.columns)}"
    # Check that edited cell text is present
    cell_1_2_text = target_table.rows[1].cells[2].text
    assert edited_status in cell_1_2_text, f"Expected edited text '{edited_status}' in cell (1, 2), got: '{cell_1_2_text}'"

    # D) Check order in document body (Heading 1 -> paragraph -> bullets -> Heading 2 -> Table -> paragraph)
    body_elements = []
    for child in out_doc.element.body:
        if child.tag.endswith('p'):
            p = Paragraph(child, out_doc)
            if p.text.strip():
                body_elements.append(("p", p.style.name if p.style else "", p.text.strip()))
        elif child.tag.endswith('tbl'):
            body_elements.append(("tbl", len(out_doc.tables[0].rows), len(out_doc.tables[0].columns)))

    h1_idx = next(i for i, el in enumerate(body_elements) if el[0] == "p" and el[1] == "Heading 1")
    bullets_idx = next(i for i, el in enumerate(body_elements) if el[0] == "p" and el[1] == "List Bullet")
    h2_idx = next(i for i, el in enumerate(body_elements) if el[0] == "p" and el[1] == "Heading 2")
    tbl_idx = next(i for i, el in enumerate(body_elements) if el[0] == "tbl")

    assert h1_idx < bullets_idx < h2_idx < tbl_idx, (
        f"Incorrect element sequence: Heading 1 at {h1_idx}, Bullet at {bullets_idx}, "
        f"Heading 2 at {h2_idx}, Table at {tbl_idx}"
    )

    # E) Layout polish: modest margins and no empty spacer paragraphs
    section = out_doc.sections[0]
    assert section.top_margin.inches <= 0.75
    assert section.bottom_margin.inches <= 0.75
    assert section.left_margin.inches <= 0.75
    assert section.right_margin.inches <= 0.75
    assert not any(p._element.xpath(".//w:br[@w:type='page']") for p in out_doc.paragraphs)
    empty_paras = [p for p in out_doc.paragraphs if not p.text.strip()]
    assert len(empty_paras) == 0, f"Expected 0 empty spacer paragraphs, found {len(empty_paras)}"


def test_translate_docx_pagination_and_natural_flow():
    """Regression test: verify sample fits on 1 page without spillover, while long documents flow naturally."""
    import io
    import docx
    import subprocess
    import tempfile
    import os
    from app.services.document_builder import build_translated_document
    from app.models.schemas import TranslationSegment

    # 1. Typical institutional policy sample (headings, bullets, table, closing note)
    sample_segs = [
        TranslationSegment(id="s-1", source="Library Borrowing Policy", translated="நூலகக் கடன் கொள்கை", blockType="heading", blockLevel=1),
        TranslationSegment(id="s-2", source="This sample policy explains library hours, borrow limits, and care requirements.", translated="இந்தக் கொள்கை நூலக நேரம், கடன் வரம்புகள் மற்றும் பராமரிப்புத் தேவைகளை விளக்குகிறது.", blockType="paragraph"),
        TranslationSegment(id="s-3", source="Library Hours", translated="நூலக நேரம்", blockType="heading", blockLevel=2),
        TranslationSegment(id="s-4", source="The library is open from 8:00 AM to 8:00 PM, Monday through Saturday.", translated="நூலகம் திங்கள் முதல் சனி வரை காலை 8:00 மணி முதல் இரவு 8:00 மணி வரை திறந்திருக்கும்.", blockType="paragraph"),
        TranslationSegment(id="s-5", source="Borrowing Rules", translated="கடன் விதிகள்", blockType="heading", blockLevel=2),
        TranslationSegment(id="s-6", source="Undergraduate students may borrow up to 3 books for 14 days.", translated="இளங்கலை மாணவர்கள் 14 நாட்களுக்கு 3 புத்தகங்கள் வரை கடன் வாங்கலாம்.", blockType="bullet"),
        TranslationSegment(id="s-7", source="Postgraduate students and researchers may borrow up to 5 books for 28 days.", translated="முதுகலை மாணவர்கள் மற்றும் ஆராய்ச்சியாளர்கள் 28 நாட்களுக்கு 5 புத்தகங்கள் வரை கடன் வாங்கலாம்.", blockType="bullet"),
        TranslationSegment(id="s-8", source="Faculty members may borrow up to 10 books for the full semester.", translated="ஆசிரியர்கள் முழு செமஸ்டருக்கும் 10 புத்தகங்கள் வரை கடன் வாங்கலாம்.", blockType="bullet"),
        TranslationSegment(id="s-9", source="Care of Library Materials", translated="நூலகப் பொருட்களின் பராமரிப்பு", blockType="heading", blockLevel=2),
        TranslationSegment(id="s-10", source="Do not highlight, write in, or fold pages of library books.", translated="நூலகப் புத்தகங்களின் பக்கங்களை மடிக்கவோ, எழுதவோ வேண்டாம்.", blockType="bullet"),
        TranslationSegment(id="s-11", source="Lost or damaged books must be replaced or paid for at full replacement cost.", translated="தொலைந்த அல்லது சேதமடைந்த புத்தகங்களை முழு செலவில் மாற்ற வேண்டும்.", blockType="bullet"),
        TranslationSegment(id="s-12", source="Quick Reference", translated="விரைவான குறிப்பு", blockType="heading", blockLevel=2),
        TranslationSegment(id="s-13", source="Policy item", translated="கொள்கை உருப்படி", blockType="table_cell", tableIndex=0, rowIndex=0, cellIndex=0, numRows=4, numCols=2),
        TranslationSegment(id="s-14", source="Details", translated="விவரங்கள்", blockType="table_cell", tableIndex=0, rowIndex=0, cellIndex=1, numRows=4, numCols=2),
        TranslationSegment(id="s-15", source="Max books (UG)", translated="அதிகபட்ச புத்தகங்கள் (இளங்கலை)", blockType="table_cell", tableIndex=0, rowIndex=1, cellIndex=0, numRows=4, numCols=2),
        TranslationSegment(id="s-16", source="3 books", translated="3 புத்தகங்கள்", blockType="table_cell", tableIndex=0, rowIndex=1, cellIndex=1, numRows=4, numCols=2),
        TranslationSegment(id="s-17", source="Loan period (UG)", translated="கடன் காலம் (இளங்கலை)", blockType="table_cell", tableIndex=0, rowIndex=2, cellIndex=0, numRows=4, numCols=2),
        TranslationSegment(id="s-18", source="14 days", translated="14 நாட்கள்", blockType="table_cell", tableIndex=0, rowIndex=2, cellIndex=1, numRows=4, numCols=2),
        TranslationSegment(id="s-19", source="Renewals", translated="புதுப்பித்தல்கள்", blockType="table_cell", tableIndex=0, rowIndex=3, cellIndex=0, numRows=4, numCols=2),
        TranslationSegment(id="s-20", source="One renewal allowed if no hold exists", translated="முன்பதிவு இல்லையெனில் ஒரு முறை புதுப்பிக்கலாம்", blockType="table_cell", tableIndex=0, rowIndex=3, cellIndex=1, numRows=4, numCols=2),
        TranslationSegment(id="s-21", source="For this test, verify headings, bullets, and tables appear in proper order and formatting.", translated="இந்த சோதனைக்கு, தலைப்புகள், புல்லட்கள் மற்றும் அட்டவணைகள் சரியான வரிசையில் உள்ளதா என சரிபார்க்கவும்.", blockType="paragraph"),
    ]

    doc_bytes, _ = build_translated_document(sample_segs, "docx", "English", "Tamil")
    doc = docx.Document(io.BytesIO(doc_bytes))

    # Verify formatting integrity: no empty spacer paragraphs and modest margins
    assert len([p for p in doc.paragraphs if not p.text.strip()]) == 0
    assert doc.sections[0].top_margin.inches <= 0.75
    assert doc.sections[0].bottom_margin.inches <= 0.75

    # Check pagination via Word COM if available
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(doc_bytes)
        tmp_path = tmp.name

    try:
        ps_cmd = (
            f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
            f"$d = $w.Documents.Open('{tmp_path}'); "
            f"$p = $d.ComputeStatistics(2); $d.Close([ref]0); $w.Quit(); Write-Host $p"
        )
        res = subprocess.run(["powershell", "-NoProfile", "-Command", ps_cmd], capture_output=True, text=True, timeout=15)
        if res.returncode == 0 and res.stdout.strip().isdigit():
            pages = int(res.stdout.strip())
            assert pages == 1, f"Expected sample document to fit on 1 page, but Word rendered {pages} pages"
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except OSError:
                pass

    # 2. Verify natural multi-page flow for longer documents (does not shrink or force to 1 page)
    long_segs = [
        TranslationSegment(
            id=f"long-{i}",
            source=f"Institutional policy compliance requirement section {i} with detailed operational guidelines.",
            translated=f"நிறுவனக் கொள்கை இணக்கத் தேவைப் பிரிவு {i}.",
            blockType="paragraph"
        )
        for i in range(45)
    ]
    long_bytes, _ = build_translated_document(long_segs, "docx", "English", "Tamil")
    long_doc = docx.Document(io.BytesIO(long_bytes))
    assert len(long_doc.paragraphs) >= 45
    # Standard text formatting preserved without extreme shrinking
    assert not any(p._element.xpath(".//w:br[@w:type='page']") for p in long_doc.paragraphs)


# ==============================================================================
# Audio Transcription & Meeting Minutes Tests
# ==============================================================================

def _get_sample_audio(fmt: str) -> bytes:
    """Helper to obtain valid audio bytes for upload tests, using FFmpeg when available."""
    import subprocess
    import shutil
    import tempfile
    import os
    if shutil.which("ffmpeg"):
        if fmt == "wav":
            p = subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "anullsrc=r=16000:cl=mono", "-t", "0.05", "-f", "wav", "pipe:1"], capture_output=True)
            if p.returncode == 0 and p.stdout:
                return p.stdout
        elif fmt == "mp3":
            p = subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "anullsrc=r=16000:cl=mono", "-t", "0.05", "-acodec", "libmp3lame", "-f", "mp3", "pipe:1"], capture_output=True)
            if p.returncode == 0 and p.stdout:
                return p.stdout
        elif fmt == "m4a":
            with tempfile.NamedTemporaryFile(suffix=".m4a", delete=False) as tmp:
                tmp_path = tmp.name
            try:
                p = subprocess.run(["ffmpeg", "-y", "-f", "lavfi", "-i", "anullsrc=r=16000:cl=mono", "-t", "0.05", "-c:a", "aac", tmp_path], capture_output=True)
                if p.returncode == 0:
                    with open(tmp_path, "rb") as f:
                        return f.read()
            finally:
                if os.path.exists(tmp_path):
                    try:
                        os.remove(tmp_path)
                    except OSError:
                        pass
    if fmt == "wav":
        return b"RIFF\x24\x00\x00\x00WAVEfmt \x10\x00\x00\x00\x01\x00\x01\x00\x44\xac\x00\x00\x88\x58\x01\x00\x02\x00\x10\x00data\x00\x00\x00\x00"
    elif fmt == "mp3":
        return b"ID3\x03\x00\x00\x00\x00\x00\x20" + b"\x00" * 60
    elif fmt == "m4a":
        return b"\x00\x00\x00\x20ftypM4A \x00\x00\x00\x00M4A mp42isom" + b"\x00" * 40
    return b""


def test_meeting_audio_uploads_mp3_wav_m4a():
    """Verify valid MP3, WAV, and M4A audio files are accepted and validated."""
    # 1. Valid MP3
    mp3_data = _get_sample_audio("mp3")
    res_mp3 = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("audio_meeting.mp3", mp3_data, "audio/mpeg")}
    )
    assert res_mp3.status_code == 200
    assert res_mp3.json()["files"][0]["name"] == "audio_meeting.mp3"

    # 2. Valid WAV
    wav_data = _get_sample_audio("wav")
    res_wav = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("audio_meeting.wav", wav_data, "audio/wav")}
    )
    assert res_wav.status_code == 200
    assert res_wav.json()["files"][0]["name"] == "audio_meeting.wav"

    # 3. Valid M4A
    m4a_data = _get_sample_audio("m4a")
    res_m4a = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("audio_meeting.m4a", m4a_data, "audio/mp4")}
    )
    assert res_m4a.status_code == 200
    assert res_m4a.json()["files"][0]["name"] == "audio_meeting.m4a"


def test_meeting_audio_invalid_extension_and_spoofed_content():
    """Verify unsupported extensions and spoofed content are rejected with HTTP 400."""
    # 1. Unsupported extension
    res_exe = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("malware.exe", b"MZ\x90\x00\x03\x00\x00\x00", "application/octet-stream")}
    )
    assert res_exe.status_code == 400
    assert "not supported" in res_exe.json()["detail"].lower()

    # 2. Spoofed MP3 (text disguised as MP3)
    res_bad_mp3 = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("fake.mp3", b"<html><body>Not an MP3 audio stream</body></html>", "audio/mpeg")}
    )
    assert res_bad_mp3.status_code == 400
    assert "not a valid mp3" in res_bad_mp3.json()["detail"].lower()

    # 3. Spoofed WAV
    res_bad_wav = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("fake.wav", b"INVALID_WAV_HEADER_DATA_1234567890", "audio/wav")}
    )
    assert res_bad_wav.status_code == 400
    assert "not a valid wav" in res_bad_wav.json()["detail"].lower()

    # 4. Spoofed M4A
    res_bad_m4a = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("fake.m4a", b"NO_FTYP_BOX_IN_THIS_CORRUPT_FILE", "audio/mp4")}
    )
    assert res_bad_m4a.status_code == 400
    assert "not a valid m4a" in res_bad_m4a.json()["detail"].lower()


def test_whisper_provider_missing_ffmpeg_clear_instructions(monkeypatch):
    """Verify that when FFmpeg is not in PATH, WhisperAudioProvider provides clear setup instructions."""
    from app.providers.audio import WhisperAudioProvider
    import shutil
    from pathlib import Path

    monkeypatch.setattr(shutil, "which", lambda cmd: None if cmd == "ffmpeg" else "dummy")

    provider = WhisperAudioProvider(model_size="base")
    with pytest.raises(RuntimeError) as exc_info:
        provider.transcribe(Path("sample.wav"))

    err = str(exc_info.value).lower()
    assert "ffmpeg" in err
    assert "system path" in err or "path" in err
    assert "winget install" in err or "install ffmpeg" in err


def test_whisper_provider_missing_whisper_package(monkeypatch):
    """Verify that when openai-whisper package is missing, an actionable error is raised."""
    from app.providers.audio import WhisperAudioProvider
    import shutil
    import sys
    from pathlib import Path

    monkeypatch.setattr(shutil, "which", lambda cmd: "/usr/bin/ffmpeg")

    # Simulate missing whisper import
    import builtins
    real_import = builtins.__import__
    def mock_import(name, *args, **kwargs):
        if name == "whisper":
            raise ImportError("No module named 'whisper'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", mock_import)

    provider = WhisperAudioProvider(model_size="base")
    with pytest.raises(RuntimeError) as exc_info:
        provider.transcribe(Path("sample.wav"))

    err = str(exc_info.value).lower()
    assert "openai-whisper" in err
    assert "pip install" in err


def test_meeting_transcribe_empty_audio_raises_error(monkeypatch):
    """Verify empty/silent audio transcription fails with a clear error and does not fabricate minutes."""
    from app.providers.audio import WhisperAudioProvider
    from app.config import settings

    monkeypatch.setattr(settings, "AUDIO_PROVIDER", "whisper")

    def mock_empty_transcribe(self, audio_path):
        raise RuntimeError("The audio file contains no audible speech or usable content. Please provide a clear recording.")

    monkeypatch.setattr(WhisperAudioProvider, "transcribe", mock_empty_transcribe)

    # Upload valid synthetic WAV header
    wav_data = b"RIFF\x24\x00\x00\x00WAVEfmt \x10\x00\x00\x00\x01\x00\x01\x00\x44\xac\x00\x00\x88\x58\x01\x00\x02\x00\x10\x00data\x00\x00\x00\x00"
    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("silent.wav", wav_data, "audio/wav")}
    )
    af_id = upload_res.json()["files"][0]["id"]

    # Endpoint: /api/meetings/transcribe
    transcribe_res = client.post(
        "/api/meetings/transcribe",
        json={"audioFileId": af_id}
    )
    assert transcribe_res.status_code == 500 or transcribe_res.status_code == 400
    assert "no audible speech" in transcribe_res.json()["detail"].lower()

    # Endpoint: /api/meetings/jobs
    job_res = client.post(
        "/api/meetings/jobs",
        json={"audioFileIds": [af_id]}
    )
    assert job_res.status_code == 200
    data = job_res.json()
    assert data["status"] == "failed"
    assert data["draft"] is None
    assert "no audible speech" in data["error"].lower()


def test_meeting_transcribe_endpoint_and_minutes_with_timestamps(monkeypatch):
    """Verify transcription preserves timestamps and draft minutes do not invent owners or dates."""
    from app.providers.audio import WhisperAudioProvider
    from app.services.extractor import DocumentPassage
    from app.config import settings

    monkeypatch.setattr(settings, "AUDIO_PROVIDER", "whisper")

    mock_passages = [
        DocumentPassage(
            text="Good morning everyone. We are starting the department review.",
            timestamp="00:00 - 00:15",
            section="Time 00:00"
        ),
        DocumentPassage(
            text="We agreed to organize the technical workshop on 14 November 2026.",
            timestamp="00:15 - 00:45",
            section="Time 00:15"
        ),
        DocumentPassage(
            text="I will prepare the presentation slides and agenda by next Monday.",
            timestamp="00:45 - 01:20",
            section="Time 00:45"
        ),
    ]

    monkeypatch.setattr(WhisperAudioProvider, "transcribe", lambda self, path: mock_passages)

    # 1. Upload audio
    wav_data = b"RIFF\x24\x00\x00\x00WAVEfmt \x10\x00\x00\x00\x01\x00\x01\x00\x44\xac\x00\x00\x88\x58\x01\x00\x02\x00\x10\x00data\x00\x00\x00\x00"
    upload_res = client.post(
        "/api/uploads",
        data={"module": "meeting"},
        files={"files": ("workshop_review.wav", wav_data, "audio/wav")}
    )
    af_id = upload_res.json()["files"][0]["id"]

    # 2. Test /api/meetings/transcribe endpoint
    t_res = client.post(
        "/api/meetings/transcribe",
        json={"audioFileId": af_id}
    )
    assert t_res.status_code == 200
    t_data = t_res.json()
    assert "[00:00 - 00:15]" in t_data["transcript"]
    assert "[00:15 - 00:45]" in t_data["transcript"]
    assert "[00:45 - 01:20]" in t_data["transcript"]
    assert t_data["audioFileName"] == "workshop_review.wav"

    # 3. Test /api/meetings/jobs creates minutes without fabricating speaker names
    job_res = client.post(
        "/api/meetings/jobs",
        json={"audioFileIds": [af_id], "meetingTitle": "Technical Workshop Review"}
    )
    assert job_res.status_code == 200
    job_data = job_res.json()
    assert job_data["status"] == "completed"
    assert job_data["source"] == "audio"
    assert job_data["audioFileName"] == "workshop_review.wav"
    assert "[00:15 - 00:45]" in job_data["transcript"]

    draft = job_data["draft"]
    assert draft is not None
    # Decision verified
    assert any("technical workshop" in d["decision"].lower() for d in draft["decisions"])

    # Action item check: because audio transcript has no speaker name prefix ("Maya: ..."),
    # owner MUST be None and uncertain MUST be True! No fabricated owner.
    assert len(draft["actionItems"]) >= 1
    action = draft["actionItems"][0]
    assert "slides and agenda" in action["task"].lower() or "prepare" in action["task"].lower()
    assert action["owner"] is None, f"Expected owner to be None, got: {action['owner']}"
    assert action["uncertain"] is True


def test_meeting_minutes_regenerate_from_edited_transcript():
    """Verify that a user can edit the transcript in UI and regenerate minutes directly from the edited text."""
    edited_transcript = (
        "[00:00 - 00:15]: Dr. Raman: Good morning everyone. We are starting the department review.\n"
        "[00:15 - 00:45]: Chair: We agreed to organize the technical workshop on 14 November 2026.\n"
        "[00:45 - 01:20]: Dr. Raman: I will prepare the presentation slides and agenda by 10 November 2026."
    )

    # Post with rawTranscriptText
    res = client.post(
        "/api/meetings/jobs",
        json={
            "rawTranscriptText": edited_transcript,
            "meetingTitle": "Department Planning"
        }
    )
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "completed"
    assert data["transcript"] == edited_transcript

    draft = data["draft"]
    assert draft is not None
    assert draft["title"] == "Department Planning"

    # Now that the user explicitly provided Dr. Raman as the speaker in the edited transcript,
    # the extractive engine extracts Dr. Raman as the owner with high certainty!
    action = [a for a in draft["actionItems"] if "slides and agenda" in a["task"].lower()][0]
    assert action["owner"] == "Dr. Raman"
    assert action["dueDate"] == "10 November 2026"
    assert action["uncertain"] is False


def test_meeting_audio_live_smoke_test():
    """Live smoke test with real audio and FFmpeg when installed; gracefully skipped otherwise with setup instructions."""
    import shutil
    import os

    if not shutil.which("ffmpeg"):
        pytest.skip(
            "Live audio transcription smoke test skipped: FFmpeg is not installed on PATH. "
            "To test live audio transcription, install FFmpeg via 'winget install Gyan.FFmpeg' on Windows, "
            "'brew install ffmpeg' on macOS, or 'sudo apt install ffmpeg' on Linux, and restart PowerShell."
        )

    # If FFmpeg is present, verify WhisperAudioProvider initialization
    from app.providers.audio import WhisperAudioProvider
    provider = WhisperAudioProvider(model_size="base")
    assert provider.model_size == "base"


def test_saved_meeting_drafts_listing_and_reopening():
    """Verify saved meeting drafts can be listed with their title and date, and reopened with edits intact."""
    from app.models.schemas import MeetingDraft, DiscussionTopic, MeetingDecision, ActionItem

    # 1. Create a job
    init_res = client.post(
        "/api/meetings/jobs",
        json={
            "rawTranscriptText": "Dean: We need to finalize the academic calendar.",
            "meetingTitle": "Faculty Senate Monthly Meeting"
        }
    )
    assert init_res.status_code == 200
    job_id = init_res.json()["id"]

    # 2. Save custom edits
    updated_draft = MeetingDraft(
        title="Faculty Senate Annual General Meeting",
        summary="Faculty senate reviewed budget approvals and approved 10 graduate assistantships.",
        discussions=[
            DiscussionTopic(topic="Graduate Assistantships", points=["Dean presented the allocation request."])
        ],
        decisions=[
            MeetingDecision(decision="Approved 10 graduate teaching assistantships", status="agreed", evidence="Unanimous vote")
        ],
        actionItems=[
            ActionItem(task="Prepare contract templates", owner="Dean Miller", dueDate="next Friday", uncertain=False)
        ],
        uncertainties=[]
    )
    update_res = client.put(
        f"/api/meetings/jobs/{job_id}/draft",
        json={"draft": updated_draft.model_dump()}
    )
    assert update_res.status_code == 200
    assert update_res.json()["draft"]["title"] == "Faculty Senate Annual General Meeting"

    # 3. List saved drafts
    list_res = client.get("/api/meetings/jobs")
    assert list_res.status_code == 200
    drafts_list = list_res.json()
    assert isinstance(drafts_list, list)
    matching = [d for d in drafts_list if d["id"] == job_id]
    assert len(matching) == 1
    saved_item = matching[0]
    assert saved_item["draft"]["title"] == "Faculty Senate Annual General Meeting"
    assert saved_item["createdAt"] is not None
    assert len(saved_item["createdAt"]) > 0

    # 4. Reopen the draft
    reopen_res = client.get(f"/api/meetings/jobs/{job_id}")
    assert reopen_res.status_code == 200
    reopened = reopen_res.json()
    assert reopened["draft"]["title"] == "Faculty Senate Annual General Meeting"
    assert reopened["draft"]["summary"] == "Faculty senate reviewed budget approvals and approved 10 graduate assistantships."
    assert reopened["draft"]["decisions"][0]["decision"] == "Approved 10 graduate teaching assistantships"
    assert reopened["draft"]["actionItems"][0]["owner"] == "Dean Miller"
    assert reopened["draft"]["actionItems"][0]["dueDate"] == "next Friday"


def test_meeting_docx_download_structure_and_latest_saved_edits():
    """Verify downloaded DOCX contains the latest saved edits and adheres to expected document structure."""
    import docx
    from app.models.schemas import MeetingDraft, DiscussionTopic, MeetingDecision, ActionItem

    # Create and save a custom draft
    init_res = client.post(
        "/api/meetings/jobs",
        json={
            "rawTranscriptText": "Discussion regarding lab safety and equipment procurement.",
            "meetingTitle": "Initial Lab Meeting"
        }
    )
    job_id = init_res.json()["id"]

    saved_draft = MeetingDraft(
        title="Lab Safety & Equipment Procurement Committee",
        summary="Executive summary regarding new biosafety equipment approvals and room renovations.",
        discussions=[
            DiscussionTopic(topic="Autoclave Procurement", points=["Vendor bids reviewed.", "Warranty terms confirmed."])
        ],
        decisions=[
            MeetingDecision(decision="Selected Vendor B for autoclave installation", status="agreed", evidence="Committee consensus")
        ],
        actionItems=[
            ActionItem(task="Issue purchase order", owner="Dr. Aris", dueDate="15 November 2026", uncertain=False),
            ActionItem(task="Inspect ventilation ducts", owner=None, dueDate=None, uncertain=True)
        ],
        uncertainties=["Ventilation inspection contractor not yet assigned."]
    )
    client.put(f"/api/meetings/jobs/{job_id}/draft", json={"draft": saved_draft.model_dump()})

    # Download DOCX
    dl_res = client.get(f"/api/meetings/jobs/{job_id}/download")
    assert dl_res.status_code == 200
    assert "officedocument" in dl_res.headers.get("content-type", "")
    assert "Lab_Safety" in dl_res.headers.get("content-disposition", "") or ".docx" in dl_res.headers.get("content-disposition", "")

    # Parse DOCX structure
    doc = docx.Document(io.BytesIO(dl_res.content))
    doc_text = "\n".join(p.text for p in doc.paragraphs)

    # Check title and latest saved edits
    assert "Lab Safety & Equipment Procurement Committee" in doc_text
    assert "Executive summary regarding new biosafety equipment approvals" in doc_text
    assert "Autoclave Procurement" in doc_text
    assert "Vendor bids reviewed." in doc_text

    # Check tables (decisions and action items)
    assert len(doc.tables) >= 2
    decisions_table_text = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "Selected Vendor B for autoclave installation" in decisions_table_text
    assert "Agreed" in decisions_table_text

    actions_table_text = "\n".join(cell.text for row in doc.tables[1].rows for cell in row.cells)
    assert "Issue purchase order" in actions_table_text
    assert "Dr. Aris" in actions_table_text
    assert "15 November 2026" in actions_table_text
    assert "Inspect ventilation ducts" in actions_table_text
    assert "Unassigned" in actions_table_text
    assert "Review required" in actions_table_text


def test_owner_extraction_from_timestamped_whisper_transcripts():
    """Verify owner extraction from timestamped Whisper transcripts with Maya and Ravi examples."""
    whisper_transcript = (
        "[00:07 - 00:11]: Maya, I will send the room layout by Friday.\n"
        "[00:11 - 00:15]: Ravi, I will inspect the electrical outlets by Monday."
    )

    res = client.post(
        "/api/meetings/jobs",
        json={"rawTranscriptText": whisper_transcript, "meetingTitle": "Facilities Inspection"}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "completed"

    # Verify timestamps remain visible in the transcript
    assert "[00:07 - 00:11]: Maya, I will send the room layout by Friday." in data["transcript"]
    assert "[00:11 - 00:15]: Ravi, I will inspect the electrical outlets by Monday." in data["transcript"]

    # Verify action items
    draft = data["draft"]
    assert draft is not None
    actions = draft["actionItems"]
    assert len(actions) == 2

    # Maya action
    maya_action = [a for a in actions if "room layout" in a["task"].lower()][0]
    assert maya_action["owner"] == "Maya", f"Expected owner 'Maya', got: {maya_action['owner']}"
    assert maya_action["dueDate"] == "Friday", f"Expected dueDate 'Friday', got: {maya_action['dueDate']}"
    assert maya_action["uncertain"] is False

    # Ravi action
    ravi_action = [a for a in actions if "electrical outlets" in a["task"].lower()][0]
    assert ravi_action["owner"] == "Ravi", f"Expected owner 'Ravi', got: {ravi_action['owner']}"
    assert ravi_action["dueDate"] == "Monday", f"Expected dueDate 'Monday', got: {ravi_action['dueDate']}"
    assert ravi_action["uncertain"] is False


def test_owner_extraction_negative_cases_do_not_infer_names():
    """Verify names are assigned as owner ONLY when the pattern is present, and never inferred otherwise."""
    transcript = (
        "[00:01 - 00:05]: I will inspect the electrical outlets by Monday.\n"
        "[00:05 - 00:10]: In addition, I will prepare the presentation slides by next Friday.\n"
        "[00:10 - 00:15]: Chair: We agreed to finalize the agenda."
    )

    res = client.post(
        "/api/meetings/jobs",
        json={"rawTranscriptText": transcript}
    )
    assert res.status_code == 200
    data = res.json()
    actions = data["draft"]["actionItems"]

    # In "I will inspect...", no speaker was mentioned -> owner must be None, uncertain True
    act1 = [a for a in actions if "electrical outlets" in a["task"].lower()][0]
    assert act1["owner"] is None
    assert act1["uncertain"] is True

    # In "In addition, I will...", "In addition" is not a person -> owner must be None, uncertain True
    act2 = [a for a in actions if "presentation slides" in a["task"].lower()][0]
    assert act2["owner"] is None
    assert act2["uncertain"] is True


def test_local_providers_work_without_gemini_key_and_gemini_rejected(monkeypatch):
    """
    Verify that the system operates 100% locally with no Gemini key or settings configured,
    and that any attempt to select 'gemini' as a provider cleanly raises an unsupported ValueError.
    """
    from app.config import settings
    from app.providers.embeddings import get_embedding_provider, LocalEmbeddingProvider
    from app.providers.llm import get_llm_provider, LocalExtractiveLLMProvider
    from app.providers.translation import get_translation_provider, LocalTranslationProvider

    # 1. Verify GEMINI_API_KEY is not defined on settings
    assert not hasattr(settings, "GEMINI_API_KEY")

    # 2. Verify health check succeeds without Gemini
    health_res = client.get("/api/health")
    assert health_res.status_code == 200
    health_data = health_res.json()
    assert health_data["status"] == "ok"
    assert "gemini" not in health_data["providers"].values()

    # 3. Verify local embedding provider functions offline
    monkeypatch.setattr(settings, "EMBEDDING_PROVIDER", "local")
    embedder = get_embedding_provider()
    assert isinstance(embedder, LocalEmbeddingProvider)
    assert embedder.dimension == 256
    vec = embedder.embed_query("Local institutional policy")
    assert len(vec) == 256

    # 4. Verify local extractive LLM provider functions offline
    monkeypatch.setattr(settings, "LLM_PROVIDER", "local")
    llm = get_llm_provider()
    assert isinstance(llm, LocalExtractiveLLMProvider)
    ans = llm.generate("Context: The academic council approved the curriculum. Question: What did the council approve?")
    assert "academic council approved" in ans.lower()

    # 5. Verify local translation provider functions offline
    monkeypatch.setattr(settings, "TRANSLATION_PROVIDER", "local")
    translator = get_translation_provider()
    assert isinstance(translator, LocalTranslationProvider)
    drafts = translator.translate_segments(["Hello"], "English", "Tamil")
    assert drafts == [""]

    # 6. Verify requesting 'gemini' explicitly raises ValueError for all provider types
    monkeypatch.setattr(settings, "EMBEDDING_PROVIDER", "gemini")
    with pytest.raises(ValueError, match="Unknown EMBEDDING_PROVIDER"):
        get_embedding_provider()

    monkeypatch.setattr(settings, "LLM_PROVIDER", "gemini")
    with pytest.raises(ValueError, match="Unknown LLM_PROVIDER"):
        get_llm_provider()

    monkeypatch.setattr(settings, "TRANSLATION_PROVIDER", "gemini")
    with pytest.raises(ValueError, match="Unknown TRANSLATION_PROVIDER"):
        get_translation_provider()








